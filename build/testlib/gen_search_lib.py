# Test libraries for the full-text search (build/scripts/search-check.js), Python with PyMuPDF, PIL and python-docx.
#
#   python build/testlib/gen_search_lib.py small <folder>
#       a dozen books with known words on known pages: PDFs (German and English, a word hyphenated at a line end,
#       umlauts written both ways), an EPUB with chapters and a page list, an EPUB without one, TXT, Markdown,
#       HTML, DOCX, RTF, a scanned PDF without a text layer, and aside in "$extra" a long PDF of 9600 pages (a new
#       book during the check, pause and resume tests).
#       <folder>/search-truth.json says where every test word stands.
#   python build/testlib/gen_search_lib.py large <folder> <source library> [scale]
#       the folders, file names and page counts of <source library> (e.g. books2\lib-master, whose PDFs carry
#       little text) with about 2000 characters of generated German or English prose on every page; `scale` (0..1)
#       keeps that share of each book's pages. A few planted words on known pages go to <folder>/search-truth.json.
import io, json, os, random, sys, zipfile, zlib

import fitz
from PIL import Image, ImageDraw, ImageFont

FONT_DIR = r'C:\Windows\Fonts'

# ------------------------------------------------------------------------------------------------ prose
DE_WORDS = ('der die das und ist nicht ein eine mit von zu den im auf für sich auch es dem als an nach wie aus bei '
            'so noch nur oder aber wenn sein war werden hat wird sie er ihr wir man schon dann über unter vor durch '
            'Zeit Jahr Welt Leben Mensch Haus Stadt Land Frage Arbeit Spiel Buch Geschichte Teil Weg Hand Auge Wasser '
            'Kraft Bild Wort Licht Abend Morgen Nacht Tag Garten Fenster Straße Brief Freund Kind Mutter Vater Schule '
            'Wissen Sprache Gedanke Ordnung Freiheit Zukunft Erfahrung Beispiel Entwicklung Gesellschaft Wirtschaft '
            'Markt Preis Geld Bank Krieg Frieden Staat Macht Recht Gesetz Grenze Reise Insel Wald Berg Fluss Meer '
            'Turm Läufer Springer Bauer König Dame Zug Partie Stellung Angriff Verteidigung Eröffnung Endspiel '
            'groß klein alt neu gut lang kurz hoch tief weit nah schnell langsam leicht schwer klar dunkel hell '
            'sagen machen geben kommen gehen sehen wissen finden denken nehmen stehen liegen bleiben halten führen '
            'zeigen spielen lesen schreiben lernen suchen bringen tragen setzen fallen ziehen erklären verstehen').split()
EN_WORDS = ('the of and to in is that it was for on are as with his they at be this from have or by one had not but '
            'what all were when we there can an your which their said if do will each about how up out them then she '
            'time year world life people house city country question work game book history part way hand eye water '
            'power picture word light evening morning night day garden window street letter friend child mother father '
            'knowledge language thought order freedom future experience example development society economy market '
            'price money bank war peace state law border journey island forest mountain river sea tower bishop knight '
            'pawn king queen move position attack defence opening endgame large small old new good long short high '
            'deep wide near fast slow light heavy clear dark bright say make give come go see know find think take '
            'stand lie stay hold lead show play read write learn search bring carry set fall pull explain understand').split()
SYL_DE = 'an be bin dor er fal gen ha hei kan lan lin mar mer nor ober ran rein sach schaf sen stein tal tor un ver wald wer zei zu'.split()
SUF_DE = ['ung', 'heit', 'keit', 'schaft', 'er', 'en', 'e', 'es', 'lich', 'isch', 'tum', 'nis']
SYL_EN = 'al ber cal con der fen gar hol lan mar mor nel per ran ster ten ton ver wil win'.split()
SUF_EN = ['ing', 'ed', 's', 'er', 'ly', 'ness', 'ment', 'tion', 'al', 'ous']


def vocabulary(lang, rnd, size=9000):
    """Common words first, then made-up words with real endings: a long tail like real text."""
    base = DE_WORDS if lang == 'de' else EN_WORDS
    syl, suf = (SYL_DE, SUF_DE) if lang == 'de' else (SYL_EN, SUF_EN)
    words = list(base)
    seen = set(words)
    while len(words) < size:
        w = ''.join(rnd.choice(syl) for _ in range(rnd.randint(2, 4))) + rnd.choice(suf)
        if lang == 'de' and rnd.random() < 0.45:
            w = w.capitalize()
        if w not in seen:
            seen.add(w)
            words.append(w)
    return words


def zipf_picker(words, rnd):
    weights = [1.0 / (i + 1) ** 1.05 for i in range(len(words))]
    cum, s = [], 0.0
    for x in weights:
        s += x
        cum.append(s)
    import bisect

    def pick():
        return words[bisect.bisect_left(cum, rnd.random() * s)]
    return pick


def prose(pick, rnd, chars):
    out, n = [], 0
    while n < chars:
        k = rnd.randint(6, 18)
        sent = ' '.join(pick() for _ in range(k))
        sent = sent[0].upper() + sent[1:] + rnd.choice(['.', '.', '.', ',', '?', '!']) if sent else sent
        out.append(sent)
        n += len(sent) + 1
    return ' '.join(out)


# ------------------------------------------------------------------------------------------------ writers
def pdf_from_pages(path, pages, title='', author='', fontsize=9.5):
    """pages: a list of strings (wrapped into the page) or of lists of lines (written line by line, exact breaks)."""
    doc = fitz.open()
    for body in pages:
        page = doc.new_page(width=420, height=595)
        if isinstance(body, list):
            y = 60
            for line in body:
                page.insert_text((40, y), line, fontsize=11, fontname='helv')
                y += 15
        elif body:
            page.insert_textbox(fitz.Rect(36, 40, 384, 560), body, fontsize=fontsize, fontname='helv')
    doc.set_metadata({'title': title, 'author': author})
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path, garbage=0, deflate=True)
    doc.close()


def scanned_pdf(path, lines_per_page, title=''):
    """Pages that are pictures of text: no text layer at all."""
    doc = fitz.open()
    try:
        font = ImageFont.truetype(os.path.join(FONT_DIR, 'times.ttf'), 26)
    except OSError:
        font = ImageFont.load_default()
    for lines in lines_per_page:
        img = Image.new('L', (840, 1190), 245)
        d = ImageDraw.Draw(img)
        for i, line in enumerate(lines):
            d.text((70, 90 + i * 40), line, fill=30, font=font)
        buf = io.BytesIO()
        img.save(buf, 'JPEG', quality=70)
        page = doc.new_page(width=420, height=595)
        page.insert_image(page.rect, stream=buf.getvalue())
    doc.set_metadata({'title': title})
    doc.save(path, garbage=0, deflate=True)
    doc.close()


def epub(path, title, author, lang, chapters, page_list=True):
    """chapters: [(title, [paragraph, …])]; with page_list every 3rd paragraph starts a print page (pagebreak + nav)."""
    pages = []
    files = []
    pno = 0
    for i, (ctitle, paras) in enumerate(chapters):
        body = [f'<h2 id="ch{i + 1}">{ctitle}</h2>']
        for k, p in enumerate(paras):
            if page_list and k % 3 == 0:
                pno += 1
                body.append(f'<span epub:type="pagebreak" id="pg{pno}" title="{pno + 10}"/>')
                pages.append((f'c{i + 1}.xhtml#pg{pno}', str(pno + 10)))
            body.append(f'<p>{p}</p>')
        files.append((f'c{i + 1}.xhtml', ctitle, ''.join(body)))
    xhtml = lambda t, b: ('<?xml version="1.0" encoding="utf-8"?><!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" '
                          f'xmlns:epub="http://www.idpf.org/2007/ops" lang="{lang}"><head><title>{t}</title></head><body>{b}</body></html>')
    nav = '<nav epub:type="toc"><h1>Inhalt</h1><ol>' + ''.join(f'<li><a href="{f}#ch{k + 1}">{t}</a></li>' for k, (f, t, _) in enumerate(files)) + '</ol></nav>'
    if page_list:
        nav += '<nav epub:type="page-list" hidden=""><ol>' + ''.join(f'<li><a href="{h}">{lab}</a></li>' for h, lab in pages) + '</ol></nav>'
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo('mimetype'), 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
        z.writestr('META-INF/container.xml', '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr('OEBPS/nav.xhtml', xhtml('Inhalt', nav))
        for f, t, b in files:
            z.writestr('OEBPS/' + f, xhtml(t, b))
        z.writestr('OEBPS/content.opf', '<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
                   f'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">lesesaal-test-{zlib.crc32(title.encode("utf-8"))}</dc:identifier>'
                   f'<dc:title>{title}</dc:title><dc:creator>{author}</dc:creator><dc:language>{lang}</dc:language>'
                   '<meta property="dcterms:modified">2026-10-04T00:00:00Z</meta></metadata><manifest>'
                   '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
                   + ''.join(f'<item id="c{k + 1}" href="{f}" media-type="application/xhtml+xml"/>' for k, (f, _, _) in enumerate(files))
                   + '</manifest><spine>' + ''.join(f'<itemref idref="c{k + 1}"/>' for k in range(len(files))) + '</spine></package>')
    return pages


def docx(path, title, paragraphs):
    import docx as pydocx
    d = pydocx.Document()
    d.core_properties.title = title
    d.add_heading(title, 1)
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(path)


def rtf(path, paragraphs):
    def esc(s):
        out = []
        for ch in s:
            o = ord(ch)
            if ch in '\\{}':
                out.append('\\' + ch)
            elif o < 128:
                out.append(ch)
            elif o < 256:
                out.append("\\'%02x" % o)
            else:
                out.append('\\u%d?' % (o if o < 32768 else o - 65536))
        return ''.join(out)
    body = ''.join('\\pard ' + esc(p) + '\\par\n' for p in paragraphs)
    open(path, 'w', encoding='latin1').write('{\\rtf1\\ansi\\ansicpg1252\\deff0{\\fonttbl{\\f0 Times New Roman;}}\\f0\\fs24\n' + body + '}')


# ------------------------------------------------------------------------------------------------ the small library
def small(folder):
    rnd = random.Random(1904)
    # the filler is made of made-up words only: no test word may turn up by chance
    de, en = vocabulary('de', rnd, 3000)[len(DE_WORDS):], vocabulary('en', rnd, 3000)[len(EN_WORDS):]
    pde, pen = zipf_picker(de, rnd), zipf_picker(en, rnd)
    filler = lambda n=1500, lang='de': prose(pde if lang == 'de' else pen, rnd, n)
    truth = {'books': {}, 'queries': []}

    # --- a German chess book: known words on known pages (1-based)
    pages = [['Der weiße Läufer', 'Ein Lehrbuch'], filler(), filler(), filler(), filler(), filler(), filler(), filler(), filler(), filler(), filler(), filler()]
    pages[2] = filler(700) + ' Der weiße Läufer opfert sich auf f7, und Herr Müller staunt. ' + filler(500)
    # a word hyphenated at the end of a line: "Bü-" / "cher"
    pages[4] = ['Im Regal stehen die alten Bü-', 'cher der Bibliothek, sauber', 'nach Themen geordnet. Die Ost-', 'und Westseite bleibt frei.', 'Die Schach-', 'Weltmeisterschaft beginnt.']
    pages[6] = filler(600) + ' Ein Brief von Mueller erwähnt den Quastenflosser im Aquarium. ' + filler(400)
    pages[8] = filler(500) + ' Das königliche Spiel verlangt Geduld. Der schwarze Läufer bleibt passiv. ' + filler(500)
    pages[11] = filler(500) + ' Im Endspiel entscheidet der aktive König. Eröffnungsrepertoire: Caro-Kann. ' + filler(300)
    pdf_from_pages(os.path.join(folder, 'Sachbuch', 'Der weiße Läufer - Anna Beispiel.pdf'), pages, 'Der weiße Läufer', 'Anna Beispiel')
    truth['books']['laeufer'] = 'Sachbuch/Der weiße Läufer - Anna Beispiel.pdf'

    # --- an English chess book
    pages = [['Bishop Tactics', 'A Primer'], filler(lang='en'), filler(lang='en'), filler(lang='en'), filler(lang='en'), filler(lang='en'), filler(lang='en'), filler(lang='en')]
    pages[1] = filler(600, 'en') + ' The bishop pair dominates the open board. ' + filler(400, 'en')
    pages[5] = filler(500, 'en') + ' Two stories about famous players and their quiet moves. ' + filler(400, 'en')
    pdf_from_pages(os.path.join(folder, 'Sachbuch', 'Bishop Tactics - Sam Example.pdf'), pages, 'Bishop Tactics', 'Sam Example')
    truth['books']['bishop'] = 'Sachbuch/Bishop Tactics - Sam Example.pdf'

    # --- an EPUB with chapters and a print page list (pages 11, 12, …)
    ch = []
    for i in range(4):
        ch.append((f'Kapitel {i + 1}', [filler(700) for _ in range(6)]))
    ch[1][1][2] = filler(300) + ' Am Ufer lag ein Quastenflosser, so alt wie die Bücher im Turm. ' + filler(300)
    ch[3][1][4] = filler(300) + ' Der Leuchtturm warf sein Licht über die Bucht. ' + filler(300)
    epub(os.path.join(folder, 'Romane', 'Die Reise.epub'), 'Die Reise', 'Clara Muster', 'de', ch, page_list=True)
    truth['books']['reise'] = 'Romane/Die Reise.epub'

    # --- an EPUB without page list
    ch = [(f'Teil {i + 1}', [filler(800) for _ in range(5)]) for i in range(3)]
    ch[2][1][3] = filler(300) + ' Hinter dem Hügel stand ein Windrad, das niemand kannte. ' + filler(300)
    epub(os.path.join(folder, 'Romane', 'Ohne Seiten.epub'), 'Ohne Seiten', 'Ben Muster', 'de', ch, page_list=False)
    truth['books']['ohne'] = 'Romane/Ohne Seiten.epub'

    # --- text formats
    os.makedirs(os.path.join(folder, 'Technik'), exist_ok=True)
    open(os.path.join(folder, 'Technik', 'Notizen zur Datenbank.txt'), 'w', encoding='utf-8').write(
        'Notizen zur Datenbank\n\n' + filler(1200) + '\n\nDie Volltextsuche nutzt einen invertierten Index mit Zeichenketten.\n\n' + filler(2400) + '\n\nStichwort Zitronenfalter am Ende.\n')
    open(os.path.join(folder, 'Technik', 'Anleitung.md'), 'w', encoding='utf-8').write(
        '# Anleitung\n\n' + filler(900) + '\n\n## Kapitel Zwei\n\nHier steht das **Kopfkissenbuch** im Fließtext.\n\n' + filler(900) + '\n')
    open(os.path.join(folder, 'Technik', 'Gärtnerei.html'), 'w', encoding='utf-8').write(
        '<!doctype html><html lang="de"><head><meta charset="utf-8"><title>Gärtnerei</title></head><body><h1>Gärtnerei</h1>'
        f'<p>{filler(800)}</p><p>Rosen und Tulpen blühen im Gewächshaus neben dem Kakteenbeet.</p><p>{filler(800)}</p></body></html>')
    docx(os.path.join(folder, 'Technik', 'Bericht.docx'), 'Bericht', [filler(600), 'Der Jahresbericht nennt die Wetterstation am Deich.', filler(600)])
    rtf(os.path.join(folder, 'Technik', 'Brief.rtf'), ['Liebe Grüße aus Lübeck,', filler(500), 'die Marzipantorte war köstlich.', filler(300)])
    truth['books'].update({'txt': 'Technik/Notizen zur Datenbank.txt', 'md': 'Technik/Anleitung.md', 'html': 'Technik/Gärtnerei.html',
                           'docx': 'Technik/Bericht.docx', 'rtf': 'Technik/Brief.rtf'})

    # --- a scan without a text layer
    os.makedirs(os.path.join(folder, 'Archiv'), exist_ok=True)
    scanned_pdf(os.path.join(folder, 'Archiv', 'Gescanntes Protokoll.pdf'), [['Protokoll der Sitzung', 'vom 3. März'], ['Seite zwei', 'des Protokolls']] * 3, 'Gescanntes Protokoll')
    truth['books']['scan'] = 'Archiv/Gescanntes Protokoll.pdf'

    # --- a long PDF, kept aside in "$extra" (the scanner skips folders starting with $): the check copies it into the
    # library later (a new book), pauses it while the reader is open and stops the app while it is read (resume)
    long_pages = [filler(1900) for _ in range(9600)]
    long_pages[0] = ['Das lange Buch', 'Ein Test']
    long_pages[1999] = filler(800) + ' Ganz hinten steht das Wort Pfefferminzbonbon. ' + filler(600)
    os.makedirs(os.path.join(folder, '$extra'), exist_ok=True)
    pdf_from_pages(os.path.join(folder, '$extra', 'Das lange Buch.pdf'), long_pages, 'Das lange Buch', 'Lara Lang')
    truth['books']['long'] = 'Archiv/Das lange Buch.pdf'
    truth['extra'] = {'long': '$extra/Das lange Buch.pdf'}

    truth['queries'] = [
        {'q': 'Läufer', 'book': 'laeufer', 'pages': [1, 3, 9]},
        {'q': '"weiße Läufer"', 'book': 'laeufer', 'pages': [1, 3]},
        {'q': 'Müller', 'book': 'laeufer', 'pages': [3, 7]},
        {'q': 'Mueller', 'book': 'laeufer', 'pages': [3, 7]},
        {'q': 'Buch', 'book': 'laeufer', 'books': ['laeufer', 'reise'], 'pages': [5]},
        {'q': 'Bücher', 'book': 'laeufer', 'books': ['laeufer', 'reise'], 'pages': [5]},
        {'q': 'Quastenflosser', 'book': 'reise', 'books': ['laeufer', 'reise']},
        {'q': 'Caro-Kann', 'book': 'laeufer', 'pages': [12]},
        {'q': 'bishop', 'book': 'bishop', 'pages': [1, 2]},
        {'q': 'story', 'book': 'bishop', 'pages': [6]},
        {'q': 'Leuchtturm', 'book': 'reise'},
        {'q': 'Windrad', 'book': 'ohne'},
        {'q': 'Zeichenkette', 'book': 'txt'},
        {'q': 'Zitronenfalter', 'book': 'txt'},
        {'q': 'Kopfkissenbuch', 'book': 'md'},
        {'q': 'Kakteenbeet', 'book': 'html'},
        {'q': 'Wetterstation', 'book': 'docx'},
        {'q': 'Marzipantorte', 'book': 'rtf'},
        {'q': 'Schach*', 'book': 'laeufer'},
        {'q': 'Pfefferminzbonbon', 'book': 'long', 'pages': [2000], 'later': True},
        {'q': 'Protokoll Sitzung', 'books': []},
    ]
    json.dump(truth, open(os.path.join(folder, 'search-truth.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('small library in', folder)


# ------------------------------------------------------------------------------------------------ the large library
def large(folder, source, scale=1.0):
    rnd = random.Random(2026)
    voc = {'de': vocabulary('de', rnd, 40000), 'en': vocabulary('en', rnd, 30000)}
    picks = {k: zipf_picker(v, rnd) for k, v in voc.items()}
    truth = {'planted': []}
    n_books = n_pages = 0
    for dp, dn, fn in os.walk(source):
        dn[:] = [d for d in dn if not d.startswith('$') and not d.startswith('.') and not d.endswith('-Dateien') and not d.endswith('_files')]
        for f in sorted(fn):
            src = os.path.join(dp, f)
            rel = os.path.relpath(src, source)
            dst = os.path.join(folder, rel)
            low = f.lower()
            if low.endswith('.pdf'):
                try:
                    count = fitz.open(src).page_count
                except Exception:
                    continue
                count = max(2, int(round(count * scale)))
                lang = 'en' if rnd.random() < 0.25 else 'de'
                brnd = random.Random(zlib.crc32(rel.encode('utf-8')))
                pages = [[os.path.splitext(f)[0][:60]]] + [prose(picks[lang], brnd, brnd.randint(1700, 2400)) for _ in range(count - 1)]
                if n_books % 37 == 5 and count > 20:
                    word = 'Planwort%04d' % n_books
                    p = brnd.randint(10, count - 1)
                    pages[p - 1] = pages[p - 1][:900] + f' {word} ' + pages[p - 1][900:]
                    truth['planted'].append({'q': word, 'file': rel.replace(os.sep, '/'), 'page': p})
                pdf_from_pages(dst, pages, os.path.splitext(f)[0])
                n_books += 1
                n_pages += count
            elif low.endswith(('.txt', '.md')) and not f.startswith('_'):
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                open(dst, 'w', encoding='utf-8').write(prose(picks['de'], rnd, 20000))
                n_books += 1
            if n_books and n_books % 50 == 0:
                print(n_books, 'books', n_pages, 'pages', flush=True)
    truth['books'] = n_books
    truth['pages'] = n_pages
    json.dump(truth, open(os.path.join(folder, 'search-truth.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('large library in', folder, n_books, 'books', n_pages, 'pages')


if __name__ == '__main__':
    if len(sys.argv) < 3 or sys.argv[1] not in ('small', 'large'):
        sys.exit(__doc__ if __doc__ else 'usage: gen_search_lib.py small <folder> | large <folder> <source> [scale]')
    os.makedirs(sys.argv[2], exist_ok=True)
    if sys.argv[1] == 'small':
        small(sys.argv[2])
    else:
        large(sys.argv[2], sys.argv[3], float(sys.argv[4]) if len(sys.argv) > 4 else 1.0)
